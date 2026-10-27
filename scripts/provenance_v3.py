#!/usr/bin/env python3
"""
Read v3 Excel batch output → convert to old map_provenance format → run PyMuPDF matching → output provenance.json
"""

import json, sys, os, re
from pathlib import Path
from datetime import datetime
from typing import Optional

import openpyxl
import fitz
from PIL import Image

PDF_ROOT = Path('data/private/papers')
PIPELINE_OUTPUT = Path('data/private/pipeline')
V3_BATCH = Path('data/private/extraction_batch')
MAX_SEARCH_CHARS = 120  # LCS is O(n*m²) — cap search text length

# Import matching functions from old provenance_mapper
sys.path.insert(0, 'integrations/private/extraction')
from provenance_mapper import (
    normalize, ngrams, find_in_pdf, map_provenance
)


def excel_to_records(xlsx_path: Path) -> dict:
    """Convert v3 Excel to old pipeline records format."""
    wb = openpyxl.load_workbook(str(xlsx_path), data_only=True)

    records = {"ages": [], "spots": [], "spot_ids": [], "failed": []}

    # Only age_sample_data — spots are tabular, no meaningful PDF provenance
    if "age_sample_data" in wb.sheetnames:
        ws = wb["age_sample_data"]
        headers = [ws.cell(1, c).value for c in range(1, ws.max_column + 1)]
        for r in range(2, ws.max_row + 1):
            row = {headers[c]: ws.cell(r, c + 1).value for c in range(len(headers))}
            rec = row_to_record(row)
            records["ages"].append(rec)

    wb.close()
    return records


def row_to_record(row: dict) -> dict:
    """Convert one Excel row to old pipeline record format."""
    rec = {}

    # Top-level fields
    for k in ["sample_id", "age_value", "error_1sigma", "MSWD", "n", "Is_Primary",
               "Rock_type", "Formation", "Group", "Member_Depth", "Location",
               "Region_Province", "City", "Specific_Location", "GPS", "Basin",
               "age_Ma_raw"]:
        val = row.get(k)
        if val is not None:
            if isinstance(val, str) and val.strip() == "":
                continue
            if k == "error_1sigma":
                rec["error_1s"] = val
            elif k == "age_Ma_raw":
                rec["age_Ma"] = val
            else:
                rec[k] = val

    # source_context (truncate — LCS is O(n*m²))
    sc = row.get("source_context")
    if sc:
        rec["source_context"] = sc[:MAX_SEARCH_CHARS]

    # CV dimensions: dim → {Selected_Term, Source_Context, evidence}
    cv_dims = [
        "Dating_mineral", "Isotope_system", "Dating_Technique",
        "Age_Interpretation", "Physical_Method", "Statistical_Model",
        "Distribution_State", "Error_Scale_Sourcing",
    ]
    for dim in cv_dims:
        term = row.get(dim)
        source = row.get(f"{dim}_source")
        reasoning = row.get(f"{dim}_reasoning")
        if term or source or reasoning:
            rec[dim] = {
                "Selected_Term": term or "",
                "Source_Context": (source or "")[:MAX_SEARCH_CHARS],
                "evidence": (reasoning or "")[:MAX_SEARCH_CHARS],
            }

    return rec


def find_pdf(paper_key: str) -> Optional[Path]:
    """Map v3 paper key to PDF path."""
    # v3 keys are like "hexudong2013" → need to match to PDF filename
    # We use the __ separator in xlsx name to get the Chinese title
    for f in V3_BATCH.iterdir():
        if f.suffix == ".xlsx" and f.stem.startswith(f"{paper_key}__"):
            chinese_part = f.stem.split("__", 1)[1]
            # Try to find matching PDF
            for pdf in PDF_ROOT.rglob("*.pdf"):
                if pdf.stem == chinese_part or pdf.stem.startswith(chinese_part[:30]):
                    return pdf
            # Fallback: try partial match
            for pdf in PDF_ROOT.rglob("*.pdf"):
                name_norm = re.sub(r'[^\w]', '', pdf.stem.lower())
                key_norm = re.sub(r'[^\w]', '', chinese_part.lower())
                if key_norm[:20] in name_norm or name_norm[:20] in key_norm:
                    return pdf
    return None


def run_single(xlsx_path: Path):
    stem = xlsx_path.stem
    paper_key = stem.split("__")[0]
    chinese_name = stem.split("__")[1] if "__" in stem else stem

    # Skip if already has provenance.json (any date)
    existing = list(PIPELINE_OUTPUT.glob(f"{chinese_name}_*_provenance.json"))
    if existing:
        print(f"[{paper_key}] SKIP: already has {len(existing)} provenance.json")
        return

    print(f"[{paper_key}] {chinese_name[:40]}")

    # Find PDF
    pdf_path = find_pdf(paper_key)
    if not pdf_path:
        print(f"  ERROR: PDF not found")
        return
    print(f"  PDF: {pdf_path.name}")

    # Load Excel → records
    records = excel_to_records(xlsx_path)
    total = sum(len(v) for v in records.values())
    print(f"  Records: ages={len(records['ages'])} spots={len(records['spots'])} failed={len(records['failed'])}")

    # Open PDF — limit pages (data is in first 15 pages, skip references)
    doc = fitz.open(str(pdf_path))
    max_pages = min(doc.page_count, 15)
    doc.select(list(range(max_pages)))
    page_dims = [(doc[i].rect.width, doc[i].rect.height) for i in range(max_pages)]
    print(f"  Pages: {doc.page_count} (limited to {max_pages})")

    # Map provenance
    all_provenances = []
    type_map = {"ages": "age", "failed": "failed_age", "spots": "spot", "spot_ids": "spot"}
    for key, recs in records.items():
        if not recs:
            continue
        data_type = type_map[key]
        print(f"  Mapping {key} ({len(recs)} records)...")
        all_provenances.extend(map_provenance(doc, recs, data_type))

    found = sum(1 for p in all_provenances if p["page"] is not None)
    not_found = sum(1 for p in all_provenances if p["page"] is None)
    print(f"  Found: {found}/{len(all_provenances)}")

    # Normalize bbox
    for p in all_provenances:
        if p["page"] and p["bbox"] and p["page"] <= len(page_dims):
            pw, ph = page_dims[p["page"] - 1]
            b = p["bbox"]
            p["bbox_norm"] = [b[0] / pw, b[1] / ph, b[2] / pw, b[3] / ph]

    # Find existing pages dir
    pages_dir = None
    for d in PIPELINE_OUTPUT.iterdir():
        if d.is_dir() and d.name.endswith("_pages") and chinese_name[:20] in d.name:
            pages_dir = d
            break
    if not pages_dir:
        for d in PIPELINE_OUTPUT.iterdir():
            if d.is_dir() and d.name.endswith("_pages"):
                # Fuzzy match
                d_stem = d.name.replace("_pages", "")
                # strip timestamp
                parts = d_stem.rsplit("_", 2)
                if len(parts) >= 3 and parts[-2].isdigit() and parts[-1].isdigit():
                    d_stem = "_".join(parts[:-2])
                if chinese_name[:20] in d_stem or d_stem[:20] in chinese_name:
                    pages_dir = d
                    break

    page_images = []
    if pages_dir:
        for img_file in sorted(pages_dir.iterdir()):
            if img_file.suffix == ".png":
                m = re.match(r"page_(\d+)\.png", img_file.name)
                if m:
                    try:
                        img = Image.open(img_file)
                        page_images.append({
                            "page": int(m.group(1)),
                            "image": img_file.name,
                            "width": img.size[0],
                            "height": img.size[1],
                            "dpi": 300,
                        })
                    except Exception:
                        pass

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    out_path = PIPELINE_OUTPUT / f"{chinese_name}_{timestamp}_provenance.json"

    result = {
        "paper_name": chinese_name,
        "pdf_path": str(pdf_path),
        "total_pages": doc.page_count,
        "total_fields": len(all_provenances),
        "found_count": found,
        "not_found_count": not_found,
        "page_images": page_images,
        "provenances": all_provenances,
    }

    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(result, f, ensure_ascii=False, indent=2)
    print(f"  Output: {out_path.name}")

    doc.close()


def main():
    if len(sys.argv) < 2:
        print("Usage: python provenance_v3.py <paper_key>  (e.g. hexudong2013)")
        print("       python provenance_v3.py --all")
        sys.exit(1)

    if sys.argv[1] == "--all":
        xlsx_files = sorted([f for f in V3_BATCH.glob("*.xlsx") if not f.name.startswith(".~")])
        print(f"Found {len(xlsx_files)} xlsx files")
        for xf in xlsx_files:
            run_single(xf)
            print()
    else:
        paper_key = sys.argv[1]
        xlsx_path = V3_BATCH / f"{paper_key}__{paper_key}.xlsx"
        # Find actual file
        for f in V3_BATCH.iterdir():
            if f.suffix == ".xlsx" and f.stem.startswith(f"{paper_key}__"):
                xlsx_path = f
                break
        if not xlsx_path.exists():
            print(f"ERROR: xlsx not found for {paper_key}")
            sys.exit(1)
        run_single(xlsx_path)


if __name__ == "__main__":
    main()
